from __future__ import annotations

import csv
import io
import json
from typing import Dict, List

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.models.content import Paper, Question, QuestionOption, Subject
from app.models.university import ContentMeta

OPTION_LETTERS = ["ক", "খ", "গ", "ঘ"]

LETTER_INDEX = {
    "ক": 0, "খ": 1, "গ": 2, "ঘ": 3,
    "A": 0, "B": 1, "C": 2, "D": 3,
    "a": 0, "b": 1, "c": 2, "d": 3,
    "0": 0, "1": 1, "2": 2, "3": 3,
}


def _cell(row: Dict, *names: str) -> str:
    for name in names:
        if name in row and row[name] is not None:
            value = str(row[name]).strip()
            if value:
                return value
    return ""


def parse_upload(filename: str, data: bytes) -> List[Dict]:
    lower = (filename or "").lower()
    if lower.endswith(".json"):
        payload = json.loads(data.decode("utf-8"))
        if isinstance(payload, dict) and "questions" in payload:
            payload = payload["questions"]
        return [dict(item) for item in payload]

    if lower.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [
            str(value).strip() if value is not None else "" for value in rows[0]
        ]
        result: List[Dict] = []
        for row in rows[1:]:
            values = ["" if value is None else value for value in row]
            result.append(dict(zip(headers, values)))
        return result

    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def import_rows(
    db: Session,
    paper: Paper,
    rows: List[Dict],
    default_subject_code: str = "",
) -> Dict:
    created = 0
    updated = 0
    skipped = 0
    errors: List[str] = []

    subjects = {
        subject.code: subject
        for subject in db.execute(select(Subject)).scalars().all()
    }

    for index, row in enumerate(rows, start=2):
        try:
            serial_raw = _cell(row, "serial", "question_no", "no")
            if not serial_raw:
                skipped += 1
                continue
            serial = int(float(serial_raw))

            subject_code = _cell(row, "subject_code", "subject") or default_subject_code
            subject = subjects.get(subject_code)
            if subject is None:
                errors.append(
                    "Row {row}: unknown subject '{code}'.".format(
                        row=index, code=subject_code or "-"
                    )
                )
                skipped += 1
                continue

            stem_bn = _cell(row, "stem_bn", "stem", "question_bn", "question")
            if not stem_bn:
                errors.append("Row {row}: missing question stem.".format(row=index))
                skipped += 1
                continue

            option_texts = [
                _cell(row, "option_a", "option_1", "a"),
                _cell(row, "option_b", "option_2", "b"),
                _cell(row, "option_c", "option_3", "c"),
                _cell(row, "option_d", "option_4", "d"),
            ]
            if not all(option_texts):
                options_json = _cell(row, "options")
                if options_json:
                    try:
                        option_texts = [
                            str(item) for item in json.loads(options_json)
                        ][:4]
                    except Exception:
                        pass
            if not all(option_texts):
                errors.append("Row {row}: needs 4 options.".format(row=index))
                skipped += 1
                continue

            correct_raw = _cell(row, "correct", "correct_letter", "answer")
            correct_index = LETTER_INDEX.get(correct_raw)
            if correct_index is None:
                try:
                    correct_index = int(float(correct_raw))
                except Exception:
                    correct_index = None
            if correct_index is None or not 0 <= correct_index <= 3:
                errors.append("Row {row}: invalid correct answer.".format(row=index))
                skipped += 1
                continue

            values = {
                "subject_id": subject.id,
                "chapter_bn": _cell(row, "chapter_bn", "chapter"),
                "stem_bn": stem_bn,
                "stem_en": _cell(row, "stem_en"),
                "correct_index": correct_index,
                "explanation_bn": _cell(row, "explanation_bn", "explanation"),
                "explanation_en": _cell(row, "explanation_en") or None,
                "shortcut_bn": _cell(row, "shortcut_bn", "shortcut") or None,
                "difficulty": _cell(row, "difficulty") or "medium",
            }

            existing = db.execute(
                select(Question).where(
                    Question.paper_id == paper.id,
                    Question.serial == serial,
                )
            ).scalar_one_or_none()

            if existing is None:
                question = Question(
                    university_id=paper.university_id,
                    paper_id=paper.id,
                    unit_id=paper.unit_id,
                    serial=serial,
                    **values,
                )
                db.add(question)
                db.flush()
                created += 1
            else:
                question = existing
                for key, value in values.items():
                    setattr(question, key, value)
                for option in list(question.options):
                    db.delete(option)
                db.flush()
                updated += 1

            for order, text in enumerate(option_texts):
                db.add(
                    QuestionOption(
                        question_id=question.id,
                        letter=OPTION_LETTERS[order],
                        text=text,
                        sort_order=order,
                    )
                )
        except Exception as exc:
            errors.append("Row {row}: {err}".format(row=index, err=exc))
            skipped += 1

    meta = db.execute(
        select(ContentMeta).where(
            ContentMeta.university_id == paper.university_id
        )
    ).scalar_one_or_none()
    if meta is None:
        db.add(ContentMeta(university_id=paper.university_id, content_version=1))
    else:
        meta.content_version += 1

    paper.question_count = (
        db.scalar(
            select(func.count(Question.id)).where(Question.paper_id == paper.id)
        )
        or 0
    )
    db.commit()

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }


def _bump_content_meta(db: Session, university_id: int) -> None:
    meta = db.execute(
        select(ContentMeta).where(ContentMeta.university_id == university_id)
    ).scalar_one_or_none()
    if meta is None:
        db.add(ContentMeta(university_id=university_id, content_version=1))
    else:
        meta.content_version += 1


def _refresh_paper_count(db: Session, paper: Paper) -> None:
    paper.question_count = (
        db.scalar(
            select(func.count(Question.id)).where(Question.paper_id == paper.id)
        )
        or 0
    )


def import_questions(
    db: Session,
    paper: Paper,
    items: List[Dict],
    mode: str = "upsert",
) -> Dict:
    created = 0
    updated = 0
    skipped = 0
    errors: List[str] = []

    if mode == "replace":
        db.execute(delete(Question).where(Question.paper_id == paper.id))
        db.flush()
        mode = "upsert"

    subjects = {
        subject.code: subject
        for subject in db.execute(select(Subject)).scalars().all()
    }
    max_serial = (
        db.scalar(
            select(func.max(Question.serial)).where(Question.paper_id == paper.id)
        )
        or 0
    )

    for index, item in enumerate(items, start=1):
        try:
            subject_code = str(item.get("subject_code") or "").strip()
            subject = subjects.get(subject_code)
            if subject is None:
                errors.append(
                    "Item {index}: unknown subject_code '{code}'.".format(
                        index=index, code=subject_code or "-"
                    )
                )
                skipped += 1
                continue

            options = item.get("options") or []
            if not isinstance(options, list) or len(options) < 2:
                errors.append(
                    "Item {index}: needs at least 2 options.".format(index=index)
                )
                skipped += 1
                continue

            stem_bn = str(item.get("stem_bn") or "").strip()
            images = item.get("images") or []
            if not stem_bn and not images:
                errors.append("Item {index}: empty question stem.".format(index=index))
                skipped += 1
                continue

            correct_index = item.get("correct_index")
            if correct_index is None:
                correct_letter = str(item.get("correct_letter") or "").strip()
                correct_index = LETTER_INDEX.get(correct_letter)
            if correct_index is None or not 0 <= int(correct_index) < len(options):
                errors.append(
                    "Item {index}: invalid correct answer.".format(index=index)
                )
                skipped += 1
                continue
            correct_index = int(correct_index)

            source_pk = item.get("source_pk")
            source_pk = str(source_pk) if source_pk not in (None, "") else None

            serial = item.get("serial")
            try:
                serial = int(serial) if serial not in (None, "") else None
            except (TypeError, ValueError):
                serial = None

            existing = None
            if source_pk is not None:
                existing = db.execute(
                    select(Question).where(
                        Question.paper_id == paper.id,
                        Question.source_pk == source_pk,
                    )
                ).scalar_one_or_none()
            elif serial is not None:
                existing = db.execute(
                    select(Question).where(
                        Question.paper_id == paper.id,
                        Question.serial == serial,
                    )
                ).scalar_one_or_none()

            if existing is not None and mode == "skip":
                skipped += 1
                continue

            if serial is None:
                max_serial += 1
                serial = max_serial
            elif source_pk is not None and existing is None:
                collision = db.execute(
                    select(Question.id).where(
                        Question.paper_id == paper.id,
                        Question.serial == serial,
                    )
                ).first()
                if collision is not None:
                    max_serial += 1
                    serial = max_serial

            values = {
                "subject_id": subject.id,
                "chapter_bn": str(item.get("chapter_bn") or ""),
                "stem_bn": stem_bn,
                "stem_html": item.get("stem_html"),
                "stem_en": str(item.get("stem_en") or ""),
                "correct_index": correct_index,
                "explanation_bn": str(item.get("explanation_bn") or ""),
                "explanation_html": item.get("explanation_html"),
                "shortcut_bn": item.get("shortcut_bn"),
                "difficulty": str(item.get("difficulty") or "medium"),
                "mark": item.get("mark"),
                "source": item.get("source"),
                "source_pk": source_pk,
                "images": images or None,
                "explanation_images": item.get("explanation_images") or None,
                "tags": item.get("tags"),
                "raw_json": item.get("raw_json"),
            }

            if existing is None:
                question = Question(
                    university_id=paper.university_id,
                    paper_id=paper.id,
                    unit_id=paper.unit_id,
                    serial=serial,
                    **values,
                )
                db.add(question)
                db.flush()
                created += 1
            else:
                question = existing
                if serial is not None and serial != question.serial:
                    collision = db.execute(
                        select(Question.id).where(
                            Question.paper_id == paper.id,
                            Question.serial == serial,
                            Question.id != question.id,
                        )
                    ).first()
                    if collision is None:
                        values["serial"] = serial
                for key, value in values.items():
                    setattr(question, key, value)
                for option in list(question.options):
                    db.delete(option)
                db.flush()
                updated += 1

            for order, option in enumerate(options):
                if not isinstance(option, dict):
                    option = {"text": str(option)}
                letter = str(option.get("letter") or "").strip()
                if not letter:
                    letter = OPTION_LETTERS[order] if order < 4 else str(order + 1)
                db.add(
                    QuestionOption(
                        question_id=question.id,
                        letter=letter,
                        text=str(option.get("text") or ""),
                        text_html=option.get("text_html"),
                        image_url=option.get("image_url"),
                        sort_order=order,
                    )
                )
        except Exception as exc:
            errors.append("Item {index}: {err}".format(index=index, err=exc))
            skipped += 1

    _bump_content_meta(db, paper.university_id)
    _refresh_paper_count(db, paper)
    db.commit()

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "question_count": paper.question_count,
    }
